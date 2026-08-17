#!/usr/bin/env python3
"""
依案號查出案主的年齡區間，藉此判斷家系圖要往上/往下畫幾代——
不需要使用者選、也不需要 Claude 看到實際出生年月日。

跟 fill_identity.py 一樣的隱私設計：這支腳本只回報「年齡區間 + 建議世代範圍」，
**不會印出出生年月日本身**。

年齡區間對應世代範圍（以案主為中心）：
- 18 歲以下（幼兒青少年案主）：往上兩代（案父母、案祖父母／外祖父母）
- 18–64 歲（中壯年案主）：往上下各一代（案父母、案子女）
- 65 歲以上（老年案主）：往下兩代（案子女、案孫子女／外孫子女）

用法：
    python3 classify_scope.py --case-id CASE-001 --data ../../../../明怡服務紀錄資料夾/個案資料表-demo.xlsx
    # 上面 --data 的實際值一律來自 shared/config.json 的 case_data_xlsx，
    # 這裡只是示範一個相對路徑長什麼樣子

如果案號查不到（例如案主資料不在這份 xlsx 裡、或是用互動訪談模式沒有查表基礎），
腳本會回報找不到，這時候改用逐字稿裡的年齡線索（案主自述幾歲、稱謂如「阿嬤」
「國小學生」等）由你自己判斷世代範圍即可，不用堅持一定要查表。
"""
import argparse
import re
import sys
from datetime import date

import openpyxl


def parse_dob(value):
    if hasattr(value, "year"):
        return date(value.year, value.month, value.day)
    text = str(value).strip()
    m = re.match(r"(\d{4})[/\-](\d{1,2})[/\-](\d{1,2})", text)
    if not m:
        return None
    y, mo, d = (int(g) for g in m.groups())
    return date(y, mo, d)


def classify(dob, today):
    age = today.year - dob.year - ((today.month, today.day) < (dob.month, dob.day))
    if age < 18:
        return "幼兒青少年案主", "up2", "以案主為中心，往上兩代（案父母、案祖父母／外祖父母）"
    if age < 65:
        return "中壯年案主", "updown1", "以案主為中心，往上下各一代（案父母、案子女）"
    return "老年案主", "down2", "以案主為中心，往下兩代（案子女、案孫子女／外孫子女）"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case-id", required=True)
    parser.add_argument("--data", required=True)
    parser.add_argument("--today", help="覆寫今天日期（YYYY-MM-DD），測試用")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.data, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    record = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(headers, row))
        if str(r.get("案號", "")).strip() == args.case_id:
            record = r
            break

    if record is None:
        print(f"錯誤：在資料表中找不到案號 {args.case_id}", file=sys.stderr)
        sys.exit(1)

    dob = parse_dob(record.get("出生年月日"))
    if dob is None:
        print("錯誤：出生年月日格式無法辨識，請改用逐字稿裡的年齡線索自行判斷世代範圍。", file=sys.stderr)
        sys.exit(1)

    today = date.fromisoformat(args.today) if args.today else date.today()
    label, scope, note = classify(dob, today)

    print(f"{label}｜scope={scope}｜{note}")


if __name__ == "__main__":
    main()
