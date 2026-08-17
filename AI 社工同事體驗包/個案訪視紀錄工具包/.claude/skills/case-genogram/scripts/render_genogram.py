#!/usr/bin/env python3
"""
把已確認的家庭成員/關係 JSON 畫成家系圖 PNG，並嵌入訪視紀錄 docx 的「家系圖」區塊。

符號系統採用標準家系圖規範（McGoldrick/Gerson 系統，純黑白線條，不用顏色區分
關係類型——關係類型一律靠線型/線條數量/斜線記號區分），詳見
references/notation-legend.md。

輸入 JSON 格式：
{
  "members": [
    {"id": "m1", "name": "案父", "gender": "male", "generation": 0, "deceased": true},
    {"id": "m2", "name": "案母", "gender": "female", "generation": 0},
    {"id": "m3", "name": "案主", "gender": "female", "generation": 1, "is_index": true},
    {"id": "m4", "name": "案夫", "gender": "male", "generation": 1}
  ],
  "relations": [
    {"a": "m1", "b": "m2", "type": "married"},
    {"a": "m3", "b": "m4", "type": "married", "quality": "close"},
    {"a": "m1", "b": "m3", "type": "parent"},
    {"a": "m2", "b": "m3", "type": "parent"}
  ]
}

- generation：世代編號，0 為最上層（最年長），數字越大世代越晚。
- members 的 name 請填**關係稱謂**（案主／案父／案母／案妻／案夫／案子／案女／
  案兄／案姊／案弟／案妹／案祖父／案祖母／案孫／案孫女……），不要填真實姓名，
  見 notation-legend.md 的稱謂對照表。
- relations.type（伴侶狀態，互斥，最多擇一）：married / cohabiting / separated / divorced / parent
  （手足關係不必列出，只要兩人 generation 相同且共同父母相同，畫圖時會自動並排表示）
- relations.quality（情感關係，可選，可與 type 同時疊加）：
  close / enmeshed / distant / conflict / cutoff / conflictClose
- members 可選 "twin_group": "t1" 標記雙胞胎，"identical_twin": true 標記同卵雙胞胎
  ——目前簡化為在姓名旁加註「（雙胞胎）」／「（同卵雙胞胎）」，子女連接線仍照一般
  手足畫法，不做收斂式雙胞胎連接線（畫法複雜、容易跑版，暫不實作）。

用法：
    python3 render_genogram.py --data family.json --png 家系圖.png
    python3 render_genogram.py --data family.json --png 家系圖.png \
        --embed-docx 輸出.docx --out 輸出.docx
"""
import argparse
import json
import math

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, Circle
import matplotlib.font_manager as fm

GEN_GAP = 2.2
UNIT_GAP = 1.6
COUPLE_GAP = 0.9
SYMBOL_SIZE = 0.5
LABEL_ZONE_TOP = 0.32   # 姓名文字區塊上緣（相對於符號中心往下）
LABEL_ZONE_BOTTOM = 0.95  # 姓名文字區塊下緣，親子連接線從這裡才繼續往下畫，避免疊到文字
IP_SCALE = 1.3   # 索引個案（IP）符號放大倍率

CJK_CANDIDATES = [
    "PingFang TC", "Heiti TC", "Hiragino Sans GB", "Hiragino Sans",
    "Microsoft JhengHei", "Noto Sans CJK TC", "Songti SC", "Arial Unicode MS",
]

# 親密關係（伴侶狀態，互斥，最多擇一）——全部黑色實線，
# 只靠「斜線數量」與「是否加三角形記號」區分，不使用顏色
STATUS_STYLE = {
    "married": {"slashes": 0},
    "cohabiting": {"slashes": 0, "triangle": True},   # 同居承諾關係（LT）：線上加朝下三角形
    "separated": {"slashes": 1},
    "divorced": {"slashes": 2},
}

# 情感關係——可與伴侶狀態疊加，一樣全黑，靠線型/線數/鋸齒/截斷記號區分
QUALITY_STYLE = {
    "close": {"mode": "parallel", "lines": 2},
    "enmeshed": {"mode": "parallel", "lines": 3},
    "distant": {"mode": "parallel", "lines": 1, "linestyle": ":"},
    "conflict": {"mode": "zigzag"},
    "conflictClose": {"mode": "zigzag_double"},  # 親密-敵意（矛盾關係）：雙線＋鋸齒疊加
    "cutoff": {"mode": "cutoff"},
}


def member_half_size(m):
    """該成員符號實際畫出來的半徑/半邊長（索引個案會放大）。"""
    scale = IP_SCALE if m.get("is_index") else 1.0
    return SYMBOL_SIZE * scale / 2


def pick_cjk_font():
    available = {f.name for f in fm.fontManager.ttflist}
    for name in CJK_CANDIDATES:
        if name in available:
            return name
    return None


def layout(members, relations):
    by_id = {m["id"]: m for m in members}
    gens = {}
    for m in members:
        gens.setdefault(m["generation"], []).append(m)

    # 每個世代先把「配偶／伴侶」合併成一個 unit（單身則自成一個 unit），
    # 之後移動一律以 unit 為單位整組移動，避免只搬其中一人、拆散了伴侶關係
    units_by_gen = {}
    unit_of_member = {}
    for gen, row in gens.items():
        spouse_of = {}
        for r in relations:
            if r.get("type") in STATUS_STYLE:
                a, b = r["a"], r["b"]
                if a in by_id and b in by_id and by_id[a]["generation"] == gen and by_id[b]["generation"] == gen:
                    spouse_of[a] = b
                    spouse_of[b] = a
        placed = set()
        units = []
        for m in row:
            mid = m["id"]
            if mid in placed:
                continue
            spouse_id = spouse_of.get(mid)
            if spouse_id and spouse_id not in placed and spouse_id in by_id:
                # 標準家系圖：男左女右（male on left, female on right）
                m_a, m_b = by_id[mid], by_id[spouse_id]
                if m_a.get("gender") == "female" and m_b.get("gender") == "male":
                    units.append([spouse_id, mid])  # swap so male is index 0
                else:
                    units.append([mid, spouse_id])
                placed.add(mid)
                placed.add(spouse_id)
            else:
                units.append([mid])
                placed.add(mid)
        units_by_gen[gen] = units
        for u in units:
            for mid in u:
                unit_of_member[mid] = u

    x_pos = {}
    unit_center = {}  # id(unit 內第一個成員) -> 該 unit 目前的中心 x
    for gen in sorted(units_by_gen):
        cursor = 0.0
        for unit in units_by_gen[gen]:
            if len(unit) == 2:
                x_pos[unit[0]] = cursor
                x_pos[unit[1]] = cursor + COUPLE_GAP
                center = cursor + COUPLE_GAP / 2
                cursor += COUPLE_GAP + UNIT_GAP
            else:
                x_pos[unit[0]] = cursor
                center = cursor
                cursor += UNIT_GAP
            unit_center[unit[0]] = center

    # children：x = 父母平均值；同一對父母底下多個子女依序在平均值左右展開。
    # 子女如果本身也是某個 unit（已婚／有伴侶）的一員，整個 unit 一起平移，
    # 才不會把伴侶拆散在原地。
    parents_of = {}
    for r in relations:
        if r.get("type") == "parent":
            parents_of.setdefault(r["b"], []).append(r["a"])

    children_by_parentset = {}
    for child_id, parent_ids in parents_of.items():
        key = tuple(sorted(parent_ids))
        children_by_parentset.setdefault(key, []).append(child_id)

    for parent_ids, child_ids in children_by_parentset.items():
        parent_xs = [x_pos[p] for p in parent_ids if p in x_pos]
        if not parent_xs:
            continue
        center = sum(parent_xs) / len(parent_xs)

        # 把子女依「所屬 unit」去重（同一對伴侶如果兩人剛好都是子女，只算一個位置）
        seen_units = []
        for cid in sorted(child_ids, key=lambda c: members.index(by_id[c])):
            unit = unit_of_member.get(cid, [cid])
            if unit not in seen_units:
                seen_units.append(unit)

        n = len(seen_units)
        for i, unit in enumerate(seen_units):
            target_center = center + (i - (n - 1) / 2) * UNIT_GAP
            old_center = unit_center.get(unit[0], x_pos[unit[0]])
            delta = target_center - old_center
            for mid in unit:
                x_pos[mid] += delta
            unit_center[unit[0]] = target_center

    for m in members:
        if m["id"] not in x_pos:
            x_pos[m["id"]] = 0.0

    max_gen = max(gens) if gens else 0
    y_pos = {m["id"]: (max_gen - m["generation"]) * GEN_GAP for m in members}
    return x_pos, y_pos


def _perp_unit(p1, p2):
    dx, dy = p2[0] - p1[0], p2[1] - p1[1]
    length = math.hypot(dx, dy) or 1.0
    return (-dy / length, dx / length)


def _dir_unit(p1, p2):
    dx, dy = p2[0] - p1[0], p2[1] - p1[1]
    length = math.hypot(dx, dy) or 1.0
    return (dx / length, dy / length)


def draw_parallel(ax, p1, p2, n, gap, linestyle, linewidth):
    perp = _perp_unit(p1, p2)
    total = (n - 1) * gap / 2
    for i in range(n):
        off = i * gap - total
        ox, oy = perp[0] * off, perp[1] * off
        ax.plot([p1[0] + ox, p2[0] + ox], [p1[1] + oy, p2[1] + oy],
                color="black", linestyle=linestyle, linewidth=linewidth, zorder=1)


def draw_zigzag(ax, p1, p2, n=8, amplitude=0.06, linewidth=1.3):
    perp = _perp_unit(p1, p2)
    xs = [p1[0] + (p2[0] - p1[0]) * i / n for i in range(n + 1)]
    ys = [p1[1] + (p2[1] - p1[1]) * i / n for i in range(n + 1)]
    zx, zy = [], []
    for i, (x, y) in enumerate(zip(xs, ys)):
        sign = 1 if i % 2 == 0 else -1
        edge = 0 if i in (0, n) else 1
        zx.append(x + perp[0] * amplitude * sign * edge)
        zy.append(y + perp[1] * amplitude * sign * edge)
    ax.plot(zx, zy, color="black", linewidth=linewidth, zorder=1)


def draw_cutoff_marks(ax, p1, p2, linewidth=1.3):
    ax.plot([p1[0], p2[0]], [p1[1], p2[1]], color="black", linewidth=linewidth, zorder=1)
    perp = _perp_unit(p1, p2)
    for frac in (0.42, 0.58):
        cx = p1[0] + (p2[0] - p1[0]) * frac
        cy = p1[1] + (p2[1] - p1[1]) * frac
        ax.plot([cx - perp[0] * 0.13, cx + perp[0] * 0.13],
                [cy - perp[1] * 0.13, cy + perp[1] * 0.13],
                color="black", linewidth=linewidth, zorder=2)


def draw_slash_marks(ax, p1, p2, n):
    """在線段中點畫 n 條斜線（分居=1 條、離婚=2 條）。"""
    if n <= 0:
        return
    perp = _perp_unit(p1, p2)
    dirv = _dir_unit(p1, p2)
    spacing = 0.09
    start = -(n - 1) * spacing / 2
    cx, cy = (p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2
    for i in range(n):
        off = start + i * spacing
        bx, by = cx + dirv[0] * off, cy + dirv[1] * off
        ax.plot(
            [bx - perp[0] * 0.16 - dirv[0] * 0.05, bx + perp[0] * 0.16 + dirv[0] * 0.05],
            [by - perp[1] * 0.16 - dirv[1] * 0.05, by + perp[1] * 0.16 + dirv[1] * 0.05],
            color="black", linewidth=1.3, zorder=2,
        )


def draw_triangle_marker(ax, p1, p2, size=0.09):
    """同居承諾關係：在線段中點下方畫一個朝下的小三角形。"""
    perp = _perp_unit(p1, p2)
    dirv = _dir_unit(p1, p2)
    cx, cy = (p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2
    tipx, tipy = cx - perp[0] * size * 1.6, cy - perp[1] * size * 1.6
    ax.plot(
        [cx - dirv[0] * size, tipx, cx + dirv[0] * size, cx - dirv[0] * size],
        [cy - dirv[1] * size, tipy, cy + dirv[1] * size, cy - dirv[1] * size],
        color="black", linewidth=1.2, zorder=2,
    )


def draw_status_and_quality(ax, p1, p2, status, quality):
    """畫伴侶狀態線（親密關係）與情感關係線（可疊加），情感關係線畫在伴侶狀態線外側。"""
    base_gap = 0.16
    slot = 0

    if status and status in STATUS_STYLE:
        style = STATUS_STYLE[status]
        draw_parallel(ax, p1, p2, 1, 0, "-", 1.4)
        if style.get("slashes"):
            draw_slash_marks(ax, p1, p2, style["slashes"])
        if style.get("triangle"):
            draw_triangle_marker(ax, p1, p2)
        slot = 1

    if quality and quality in QUALITY_STYLE:
        style = QUALITY_STYLE[quality]
        perp = _perp_unit(p1, p2)
        off = base_gap * (slot + 1)
        qp1 = (p1[0] + perp[0] * off, p1[1] + perp[1] * off)
        qp2 = (p2[0] + perp[0] * off, p2[1] + perp[1] * off)
        mode = style.get("mode", "parallel")
        if mode == "zigzag":
            draw_zigzag(ax, qp1, qp2)
        elif mode == "zigzag_double":
            draw_parallel(ax, qp1, qp2, 2, 0.06, "-", 1.2)
            draw_zigzag(ax, qp1, qp2)
        elif mode == "cutoff":
            draw_cutoff_marks(ax, qp1, qp2)
        else:
            draw_parallel(ax, qp1, qp2, style.get("lines", 1), 0.06, style.get("linestyle", "-"), 1.2)


def draw(members, relations, out_png):
    by_id = {m["id"]: m for m in members}
    x_pos, y_pos = layout(members, relations)

    fig, ax = plt.subplots(figsize=(max(6, len(members) * 1.4), 6))
    font_name = pick_cjk_font()
    font_kwargs = {"fontfamily": font_name} if font_name else {}

    for r in relations:
        a, b = r.get("a"), r.get("b")
        if a not in x_pos or b not in x_pos:
            continue
        p1, p2 = (x_pos[a], y_pos[a]), (x_pos[b], y_pos[b])
        rel_type = r.get("type")
        quality = r.get("quality")

        if rel_type in STATUS_STYLE or quality:
            draw_status_and_quality(ax, p1, p2, rel_type if rel_type in STATUS_STYLE else None, quality)
        elif rel_type == "parent":
            child_x, child_y = p2
            parent_x, parent_y = p1
            parent_half = member_half_size(by_id[a])
            child_half = member_half_size(by_id[b])
            # 考慮符號放大時的 label zone 調整
            scale_p = IP_SCALE if by_id[a].get("is_index") else 1.0
            label_zone_top_p = LABEL_ZONE_TOP * scale_p
            label_zone_bottom_p = LABEL_ZONE_BOTTOM * scale_p
            scale_c = IP_SCALE if by_id[b].get("is_index") else 1.0
            label_zone_top_c = LABEL_ZONE_TOP * scale_c

            mid_y = parent_y - GEN_GAP / 2
            # 線從符號下緣先畫一小段到姓名文字區塊上緣，跳過文字區塊，
            # 再從文字區塊下緣繼續往下畫到橫線，避免線疊到姓名文字
            ax.plot([parent_x, parent_x], [parent_y - parent_half, parent_y - label_zone_top_p], color="black", linewidth=1.1, zorder=1)
            ax.plot([parent_x, parent_x], [parent_y - label_zone_bottom_p, mid_y], color="black", linewidth=1.1, zorder=1)
            ax.plot([parent_x, child_x], [mid_y, mid_y], color="black", linewidth=1.1, zorder=1)
            ax.plot([child_x, child_x], [mid_y, child_y + child_half], color="black", linewidth=1.1, zorder=1)

    for m in members:
        x, y = x_pos[m["id"]], y_pos[m["id"]]
        is_index = bool(m.get("is_index"))
        is_disabled = bool(m.get("disabled"))
        scale = IP_SCALE if is_index else 1.0
        size = SYMBOL_SIZE * scale
        border_color = "black"  # 標準規範：IP 雙線框仍為黑色，不使用紅色
        border_width = 2.0 if is_index else 1.2
        border_n = 2 if is_index else 1  # 索引個案：同心雙線框

        if m.get("gender") == "female":
            for i in range(border_n):
                r = size / 2 - i * 0.05 if border_n > 1 else size / 2
                ax.add_patch(Circle((x, y), r, fill=False, edgecolor=border_color, linewidth=border_width, zorder=3))
        else:
            half = size / 2
            for i in range(border_n):
                inset = i * 0.05
                ax.add_patch(Rectangle((x - half + inset, y - half + inset), size - 2 * inset, size - 2 * inset,
                                        fill=False, edgecolor=border_color, linewidth=border_width, zorder=3))

            # 身障標記：男性方形右半邊塗黑
            if is_disabled:
                half = size / 2
                for i in range(border_n):
                    inset = i * 0.05
                    right_rect = Rectangle((x + inset, y - half + inset), half - inset, size - 2 * inset,
                                          fill=True, facecolor="black", edgecolor="none", zorder=2)
                    ax.add_patch(right_rect)

        if m.get("deceased"):
            # 已故：符號內畫滿版 X（兩條對角線）
            half = size / 2 * 0.92
            ax.plot([x - half, x + half], [y - half, y + half], color="black", linewidth=1.3, zorder=4)
            ax.plot([x - half, x + half], [y + half, y - half], color="black", linewidth=1.3, zorder=4)

        label = m.get("name", "")
        notes = []
        if m.get("note"):
            notes.append(m["note"])
        if m.get("twin_group"):
            notes.append("同卵雙胞胎" if m.get("identical_twin") else "雙胞胎")
        if notes:
            label = f"{label}（{'、'.join(notes)}）"
        label_y = y - size / 2 - (LABEL_ZONE_TOP - SYMBOL_SIZE / 2)
        ax.text(x, label_y, label, ha="center", va="top", fontsize=10, zorder=5, **font_kwargs)

    ax.set_aspect("equal")
    ax.axis("off")
    xs = list(x_pos.values()) or [0]
    ys = list(y_pos.values()) or [0]
    ax.set_xlim(min(xs) - 1.8, max(xs) + 1.8)
    ax.set_ylim(min(ys) - 1.5, max(ys) + 1.5)
    fig.tight_layout()
    fig.savefig(out_png, dpi=200, bbox_inches="tight")
    plt.close(fig)


def embed_into_docx(png_path, docx_path, out_path):
    from docx import Document
    from docx.shared import Inches

    doc = Document(docx_path)
    table = doc.tables[0]
    legend_cell = table.rows[15].cells[0]
    para = legend_cell.add_paragraph()
    run = para.add_run()
    run.add_picture(png_path, width=Inches(6))
    doc.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", required=True, help="成員/關係 JSON 路徑")
    parser.add_argument("--png", required=True, help="輸出 PNG 路徑")
    parser.add_argument("--embed-docx", help="要嵌入的訪視紀錄 docx 路徑（選填）")
    parser.add_argument("--out", help="嵌入後輸出的 docx 路徑（搭配 --embed-docx 使用）")
    args = parser.parse_args()

    with open(args.data, encoding="utf-8") as f:
        payload = json.load(f)

    draw(payload["members"], payload["relations"], args.png)
    print(f"已產出家系圖：{args.png}")

    if args.embed_docx:
        if not args.out:
            parser.error("--embed-docx 需要搭配 --out")
        embed_into_docx(args.png, args.embed_docx, args.out)
        print(f"已將家系圖嵌入 docx：{args.out}")


if __name__ == "__main__":
    main()
