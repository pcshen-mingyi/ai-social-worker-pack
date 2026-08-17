#!/usr/bin/env python3
"""
依案號從個案資料表 (xlsx) 查表，直接把身分欄位寫入訪視紀錄 docx。

設計原則（見 SKILL.md）：
- 查到的姓名／身分證字號／電話／地址等實際內容「不會印到 stdout」，
  呼叫端（Claude）只會看到成功/失敗訊息，看不到欄位內容本身。
- 只負責表格上半部的「身分欄位」，不碰下半部「訪視紀錄」敘事欄位。

用法：
    python3 fill_identity.py --case-id CASE-001 \
        --data ../../../../明怡服務紀錄資料夾/個案資料表-demo.xlsx \
        # 實際值一律來自 shared/config.json 的 case_data_xlsx，這裡只是示範
        --docx 案主草稿.docx \
        --out 案主草稿.docx   # 可與 --docx 相同，原地寫入
"""
import argparse
import re
import sys

import openpyxl
from docx import Document

FIELDS = [
    "案號", "姓名", "出生年月日", "身分證字號", "家電", "手機", "性別",
    "婚姻狀況", "戶籍地址", "現住地址", "教育程度", "居住狀況", "職業",
    "經濟狀況", "宗教信仰",
]

MARRIAGE_OPTIONS = ["未婚", "同居", "已婚", "離婚", "鰥寡", "分居", "不詳"]
EDUCATION_OPTIONS = ["不識字", "小學以下", "國小", "國中", "高中職", "大學", "研究所"]
JOB_OPTIONS = ["商", "工", "農", "漁", "公", "教", "學生", "服務業", "無", "其他"]
RELIGION_OPTIONS = ["道教", "佛教", "基督教", "天主教", "其他"]
LIVING_OPTIONS = ["居無定所", "獨居", "與配偶同住", "與子女同住", "與其他人同住"]
ECONOMY_PREFIXES = ["工作收入每月平均", "子女供給每月平均", "親友供給每月平均"]


def lookup_case(xlsx_path, case_id):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row))
        if str(record.get("案號", "")).strip() == case_id:
            return record
    return None


def get_para_text(p):
    return "".join(r.text for r in p.runs)


def set_para_text(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def set_cell_plain_text(cell, text):
    # 表格中「留白待填」的儲存格通常只有一個空段落
    para = cell.paragraphs[0]
    set_para_text(para, text)


def append_after_label(cell, label, value):
    if not value:
        return False
    for p in cell.paragraphs:
        text = get_para_text(p)
        if text.strip().rstrip("：:") == label.rstrip("：:") or text.strip() == label:
            set_para_text(p, text + str(value))
            return True
    return False


def mark_checkbox(cell, option_label, extra_text=None):
    """把儲存格內對應選項的『□』改成『■』，可選擇在選項後方補上細節文字。"""
    for p in cell.paragraphs:
        text = get_para_text(p)
        marker = "□" + option_label
        idx = text.find(marker)
        if idx == -1:
            continue
        new_text = text[:idx] + "■" + option_label + text[idx + len(marker):]
        if extra_text:
            insert_at = idx + 1 + len(option_label)
            # 範本裡有些選項後面緊接著自帶的「：」（例如「其他：」），細節文字要補在
            # 那個冒號之後，而不是插在選項文字和冒號中間
            if new_text[insert_at:insert_at + 1] in ("：", ":"):
                insert_at += 1
            new_text = new_text[:insert_at] + extra_text + new_text[insert_at:]
        set_para_text(p, new_text)
        return True
    return False


def fill_economy(cell, value):
    if not value:
        return False
    for prefix in ECONOMY_PREFIXES:
        if prefix in value:
            m = re.search(r"(\d[\d,]*)", value)
            amount = m.group(1) if m else ""
            for p in cell.paragraphs:
                text = get_para_text(p)
                marker = "□" + prefix
                idx = text.find(marker)
                if idx == -1:
                    continue
                blank_idx = text.find("＿＿＿＿", idx)
                if blank_idx == -1:
                    continue
                new_text = (
                    text[:idx] + "■" + prefix + amount
                    + text[blank_idx + len("＿＿＿＿"):]
                )
                set_para_text(p, new_text)
                return True
    return False


def fill_living(cell, value):
    if not value:
        return False
    base, _, detail = value.partition("：")
    for option in LIVING_OPTIONS:
        if option == base or (option == "與其他人同住" and base.startswith("與其他人同住")):
            # 範本裡「與其他人同住：」的冒號已經印在選項文字後面，這裡只補細節文字
            extra = detail if (option == "與其他人同住" and detail) else None
            return mark_checkbox(cell, option, extra)
    # 找不到對應選項時，誠實地不硬套，交由使用者人工確認
    return False


def fill_current_address(cell, household_address, current_address):
    if not current_address:
        return False
    if current_address.strip() == "同上":
        return mark_checkbox(cell, "同上")
    # 範本裡「其他：」的冒號已經印在選項文字後面，這裡只補地址本身
    base, _, detail = current_address.partition("：")
    if base.strip() == "其他":
        return mark_checkbox(cell, "其他", detail)
    return mark_checkbox(cell, "其他", current_address)


def fill_identity(docx_path, out_path, record):
    doc = Document(docx_path)
    table = doc.tables[0]
    rows = table.rows

    set_cell_plain_text(rows[0].cells[1], str(record.get("姓名") or ""))
    set_cell_plain_text(rows[0].cells[3], str(record.get("出生年月日") or ""))

    set_cell_plain_text(rows[1].cells[1], str(record.get("身分證字號") or ""))
    append_after_label(rows[1].cells[3], "家電：", record.get("家電"))
    append_after_label(rows[1].cells[3], "手機：", record.get("手機"))

    gender = str(record.get("性別") or "")
    if gender in ("男", "女"):
        mark_checkbox(rows[2].cells[1], gender)

    marriage = str(record.get("婚姻狀況") or "")
    if marriage in MARRIAGE_OPTIONS:
        mark_checkbox(rows[2].cells[3], marriage)

    set_cell_plain_text(rows[3].cells[1], str(record.get("戶籍地址") or ""))

    fill_current_address(
        rows[4].cells[1],
        record.get("戶籍地址"),
        str(record.get("現住地址") or ""),
    )

    education = str(record.get("教育程度") or "")
    if education in EDUCATION_OPTIONS:
        mark_checkbox(rows[5].cells[1], education)

    fill_living(rows[6].cells[1], str(record.get("居住狀況") or ""))

    job = str(record.get("職業") or "")
    if job in JOB_OPTIONS:
        mark_checkbox(rows[7].cells[1], job)
    elif job:
        mark_checkbox(rows[7].cells[1], "其他", job)

    fill_economy(rows[8].cells[1], str(record.get("經濟狀況") or ""))

    religion = str(record.get("宗教信仰") or "")
    if religion in RELIGION_OPTIONS:
        mark_checkbox(rows[9].cells[1], religion)
    elif religion:
        mark_checkbox(rows[9].cells[1], "其他", religion)

    doc.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case-id", required=True)
    parser.add_argument("--data", required=True, help="個案資料表 xlsx 路徑")
    parser.add_argument("--docx", required=True, help="來源 docx 路徑（通常是已複製好的範本）")
    parser.add_argument("--out", required=True, help="輸出 docx 路徑")
    args = parser.parse_args()

    record = lookup_case(args.data, args.case_id)
    if record is None:
        print(f"錯誤：在資料表中找不到案號 {args.case_id}", file=sys.stderr)
        sys.exit(1)

    fill_identity(args.docx, args.out, record)
    # 刻意不印出任何欄位內容，只回報成功，避免身分資料經過對話紀錄
    print(f"已成功寫入案號 {args.case_id} 的身分欄位。")


if __name__ == "__main__":
    main()
