#!/usr/bin/env python3
"""
把訪視紀錄的「敘事內容」（由 Claude 依逐字稿/筆記草擬）寫入 docx。

跟 fill_identity.py 分工不同：這支腳本處理的欄位本來就會出現在敘事文字裡
（訪視日期、地點、主述議題……），來源就是逐字稿，因此不需要對 Claude 保密。

輸入：一個 JSON 檔，欄位皆為選填，缺的欄位就保留範本原樣。
{
  "visit_number": "1",
  "visit_date": "2026/08/04",
  "visit_method": "實地",           // 電話 / 視訊 / 實地
  "visit_location": "案家",
  "visit_staff": "王社工",
  "recorder": "王社工",
  "companion": "陳督導",
  "presenting_issue": "……",
  "case_overview": "……",
  "assessment": "……",              // 可用 \n 分段
  "needs_assessment": "……",
  "short_term_plan": "……",
  "mid_long_term_plan": "……",
  "referral": "……",
  "next_follow_up": "2026/09/04"
}
"""
import argparse
import copy
import json

from docx import Document
from docx.oxml.ns import qn

METHOD_OPTIONS = ["電話", "視訊", "實地"]


def clone_run_format(new_run, reference_run):
    """複製參考 run 的完整格式（字型、大小、中文字型……）到新 run，
    避免新插入的段落跑掉範本原本的字型設定（Word 預設樣式跟範本不一致）。"""
    ref_rpr = reference_run._r.find(qn("w:rPr"))
    if ref_rpr is None:
        return
    new_rpr = copy.deepcopy(ref_rpr)
    existing = new_run._r.find(qn("w:rPr"))
    if existing is not None:
        new_run._r.remove(existing)
    new_run._r.insert(0, new_rpr)


def add_paragraph_like(anchor, cell, text, reference_run):
    """在 anchor 段落之前插入一個新段落（anchor 為 None 時改為附加在 cell 最後），
    文字格式比照 reference_run。"""
    if anchor is not None:
        new_para = anchor.insert_paragraph_before()
    else:
        new_para = cell.add_paragraph()
    new_run = new_para.add_run(text)
    clone_run_format(new_run, reference_run)
    return new_para


def get_para_text(p):
    return "".join(r.text for r in p.runs)


def set_para_text(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def find_index(cell, exact_text):
    for i, p in enumerate(cell.paragraphs):
        if get_para_text(p).strip() == exact_text:
            return i
    return None


def replace_blank(cell, para_index, label, value):
    if para_index is None or not value:
        return
    p = cell.paragraphs[para_index]
    text = get_para_text(p)
    set_para_text(p, text.replace(label + "＿＿＿＿＿＿＿＿", label + str(value), 1))


def mark_method(cell, para_index, method):
    if para_index is None or method not in METHOD_OPTIONS:
        return
    p = cell.paragraphs[para_index]
    text = get_para_text(p)
    set_para_text(p, text.replace("□" + method, "■" + method, 1))


def fill_visit_number(cell, para_index, number):
    if para_index is None or not number:
        return
    p = cell.paragraphs[para_index]
    text = get_para_text(p)
    set_para_text(p, text.replace("第　　次", f"第 {number} 次", 1))


def fill_section(cell, heading_text, content):
    if not content:
        return
    idx = find_index(cell, heading_text)
    if idx is None or idx + 1 >= len(cell.paragraphs):
        return
    lines = [ln for ln in str(content).strip().split("\n") if ln.strip()]
    if not lines:
        return
    target = cell.paragraphs[idx + 1]
    set_para_text(target, lines[0])
    reference_run = target.runs[0]
    anchor = cell.paragraphs[idx + 2] if idx + 2 < len(cell.paragraphs) else None
    for line in lines[1:]:
        add_paragraph_like(anchor, cell, line, reference_run)


def fill_plan_bullet(cell, heading_text, content):
    if not content:
        return
    idx = find_index(cell, heading_text)
    if idx is None:
        return
    lines = [ln for ln in str(content).strip().split("\n") if ln.strip()]
    paras = cell.paragraphs
    heading_para = paras[idx]
    reference_run = heading_para.runs[0] if heading_para.runs else None
    anchor = paras[idx + 1] if idx + 1 < len(paras) else None
    for line in lines:
        if reference_run is not None:
            add_paragraph_like(anchor, cell, line, reference_run)
        elif anchor is not None:
            anchor.insert_paragraph_before(line)
        else:
            cell.add_paragraph(line)


def fill_narrative(docx_path, out_path, fields):
    doc = Document(docx_path)
    table = doc.tables[0]
    rows = table.rows

    meta_cell = rows[11].cells[0]
    fill_visit_number(meta_cell, 0, fields.get("visit_number"))
    replace_blank(meta_cell, 1, "訪視日期：", fields.get("visit_date"))
    mark_method(meta_cell, 1, fields.get("visit_method"))
    replace_blank(meta_cell, 2, "訪視地點：", fields.get("visit_location"))
    replace_blank(meta_cell, 2, "訪視人員：", fields.get("visit_staff"))
    replace_blank(meta_cell, 3, "記錄人：", fields.get("recorder"))
    replace_blank(meta_cell, 3, "陪同人員：", fields.get("companion"))

    fill_section(meta_cell, "【主述議題】", fields.get("presenting_issue"))
    fill_section(meta_cell, "【個案概況】", fields.get("case_overview"))
    fill_section(meta_cell, "【各項評估】", fields.get("assessment"))
    fill_section(meta_cell, "【需求評估】", fields.get("needs_assessment"))

    plan_cell = rows[13].cells[0]
    fill_plan_bullet(plan_cell, "－ 短期目標與行動（1–3 個月）", fields.get("short_term_plan"))
    fill_plan_bullet(plan_cell, "－ 中長期目標與行動", fields.get("mid_long_term_plan"))
    fill_plan_bullet(plan_cell, "－ 轉介／資源連結建議", fields.get("referral"))
    fill_plan_bullet(plan_cell, "－ 下次追蹤時間", fields.get("next_follow_up"))

    doc.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--docx", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--json", required=True, help="敘事欄位 JSON 檔路徑")
    args = parser.parse_args()

    with open(args.json, encoding="utf-8") as f:
        fields = json.load(f)

    fill_narrative(args.docx, args.out, fields)
    print("已寫入訪視紀錄敘事內容。")


if __name__ == "__main__":
    main()
